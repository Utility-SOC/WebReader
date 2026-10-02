"""Builds a small, realistic sample repository for demos and testing (run inside the backend container).

Includes deliberate accessibility problems (untagged PDF, scan with no text layer, slide without a title,
filename used as alt text, missing document title) so the checks have something to find.
"""
import os
import random
import sys

import docx
import pptx
from PIL import Image, ImageDraw, ImageFont

from backend.tests.pdfgen import Page, build_pdf

root = sys.argv[1] if len(sys.argv) > 1 else "/repo"
random.seed(7)
for d in ("public/planning", "public/parks", "public/clerk", "public/roads", "public/drafts", "private"):
    os.makedirs(f"{root}/{d}", exist_ok=True)


def png():
    im = Image.new("RGB", (80, 50), "navy")
    import io
    b = io.BytesIO(); im.save(b, "PNG"); b.seek(0)
    return b


# Planning minutes: a complete, well-formed Word document
d = docx.Document(); d.core_properties.title = "Planning Commission Minutes - March"; d.core_properties.author = "Planning Department"
d.add_heading("Zoning amendments", 1)
d.add_paragraph("The commission heard public comment on the proposed zoning change near the river and voted to continue the item to April.")
d.save(f"{root}/public/planning/minutes-march.docx")

# A Word document with typical problems: no title, skipped heading level, picture with no alt text
d = docx.Document(); d.add_heading("Annual Report", 1); d.add_paragraph("Summary of the year's activity across all departments.")
d.add_heading("Details", 3); d.add_picture(png()); d.save(f"{root}/public/annual-report.docx")

# Budget deck: no language set, a slide without a title, a picture whose alt text is its file name
p = pptx.Presentation()
s = p.slides.add_slide(p.slide_layouts[0]); s.shapes.title.text = "County Budget Overview"; s.placeholders[1].text = "Fiscal year 2027"
s2 = p.slide_layouts[6]; s2 = p.slides.add_slide(s2)
s2.shapes.add_textbox(0, 0, 3000000, 500000).text_frame.text = "Revenue grew four percent"
s2.shapes.add_picture(png(), 0, 1000000)
p.save(f"{root}/public/budget-2027.pptx")

# Untagged text PDF
pg = Page(); pg.lines(72, 100, 11, ["Notice of public meeting regarding the library parking expansion.", "Residents may attend and comment."] * 3)
open(f"{root}/public/notice.pdf", "wb").write(build_pdf([pg], title="Library Parking Notice", lang="en"))

# A scan: an image of a page with no text layer at all (exercises the OCR fallback)
im = Image.new("L", (2550, 3300), 255); dr = ImageDraw.Draw(im)
dr.text((300, 200), "Water Rights Hearing Notice", font=ImageFont.load_default(size=84), fill=0)
for i, ln in enumerate(["The board will hear testimony about water rights", "on the north fork of the river next Tuesday evening", "in the county commission chambers.",
                        "Written comments are accepted until noon Monday."]):
    dr.text((300, 480 + i * 90), ln, font=ImageFont.load_default(size=46), fill=0)
open(f"{root}/public/scanned-hearing-notice.pdf", "wb").write(build_pdf([Page().scan(im)]))

open(f"{root}/public/snow.txt", "w").write("Snow removal priorities for county roads this winter. Main arterials are plowed first.")

topics = {"parks": ("Parks and Recreation", ["trail", "playground", "picnic", "pool", "pavilion"]),
          "clerk": ("County Clerk", ["marriage", "election", "recording", "license", "voter"]),
          "roads": ("Public Works", ["pothole", "paving", "culvert", "plowing", "bridge"])}
for folder, (agency, words) in topics.items():
    for i in range(1, 11):
        w = random.sample(words, 3)
        d = docx.Document(); d.core_properties.title = f"{agency}: {w[0].title()} Update {i}"; d.core_properties.language = "en-US"
        d.add_heading(f"{w[0].title()} update", 1)
        d.add_paragraph(f"This notice covers {w[0]}, {w[1]} and {w[2]} matters for the coming quarter. " * 3)
        d.save(f"{root}/public/{folder}/{folder}-{i:02d}.docx")

# Things that must never be published
pg = Page(); pg.lines(72, 100, 11, ["Confidential personnel matter, not for release."] * 3)
open(f"{root}/private/hr-matter.pdf", "wb").write(build_pdf([pg]))
d = docx.Document(); d.add_paragraph("Unreleased draft."); d.save(f"{root}/public/drafts/draft.docx")
print(f"Sample repository written to {root}")
