"""Structure extraction and accessibility checks on generated DOCX/PPTX files."""
import io

import pytest
from PIL import Image

docx = pytest.importorskip("docx")
pptx = pytest.importorskip("pptx")

from backend.office import process_office, process_pdf_accessibility
from backend.extract_docx import extract_docx
from backend.extract_pptx import extract_pptx
from backend import a11y_checks


def _png():
    b = io.BytesIO()
    Image.new("RGB", (60, 40), "navy").save(b, format="PNG")
    b.seek(0)
    return b


def codes(issues):
    return [i.code for i in issues]


# ---------------------------------------------------------------- DOCX

@pytest.fixture
def sample_docx(tmp_path):
    d = docx.Document()
    d.add_heading("Annual Report", 1)
    d.add_paragraph("Opening paragraph with some text.")
    d.add_heading("Details", 3)  # skips level 2
    d.add_paragraph("A bullet", style="List Bullet")
    t = d.add_table(rows=3, cols=2)
    for r, row in enumerate(t.rows):
        for c, cell in enumerate(row.cells):
            cell.text = f"r{r}c{c}"
    good = d.add_picture(_png()); good._inline.docPr.set("descr", "A navy square")
    bad = d.add_picture(_png()); bad._inline.docPr.set("descr", "image1.png")
    d.add_picture(_png())  # no alt text at all
    path = tmp_path / "a.docx"
    d.save(path)
    return str(path)


def test_docx_structure_in_reading_order(sample_docx):
    s = extract_docx(sample_docx)
    kinds = [(b.kind, b.text or b.alt_text) for b in s.blocks]
    assert kinds[0] == ("heading", "Annual Report") and s.blocks[0].level == 1
    assert s.blocks[1].kind == "paragraph"
    assert s.blocks[2].kind == "heading" and s.blocks[2].level == 3
    assert s.blocks[3].kind == "list_item"
    assert s.blocks[4].kind == "table" and s.blocks[4].rows[1] == ["r1c0", "r1c1"]
    figs = s.figures()
    assert [f.alt_text for f in figs] == ["A navy square", "image1.png", ""]


def test_docx_issues(sample_docx):
    _, acc = process_office(sample_docx, "docx")
    got = {i["code"] for i in acc["issues"]}
    assert {"DOC_TITLE_MISSING", "HEADING_SKIP", "TABLE_DESCRIPTION_MISSING", "FIG_ALT_MISSING", "FIG_ALT_PLACEHOLDER"} <= got
    assert "TABLE_NO_HEADER" not in got  # tables are treated as images by default
    # the good picture produced no figure issue
    fig_issues = [i for i in acc["issues"] if i["code"].startswith("FIG_")]
    assert len(fig_issues) == 2
    skip = next(i for i in acc["issues"] if i["code"] == "HEADING_SKIP")
    assert skip["detail"] == {"from": 1, "to": 3} and skip["wcag"] == "1.3.1"
    assert acc["summary"]["by_fix"]["suggest"] >= 3


def test_docx_reading_text_includes_alt_text_and_tables(sample_docx):
    text, _ = process_office(sample_docx, "docx")
    assert "Annual Report" in text and "Image: A navy square" in text and "Table." in text and "r1c1" in text
    assert "image1.png" in text  # present alt text is read, even though flagged


def test_docx_table_alt_text_clears_description_issue_and_is_read(tmp_path):
    d = docx.Document()
    t = d.add_table(rows=2, cols=2)
    for r in range(2):
        for c in range(2):
            t.cell(r, c).text = f"r{r}c{c}"
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pr = t._tbl.tblPr
    cap = OxmlElement("w:tblCaption"); cap.set(qn("w:val"), "Budget by quarter"); pr.append(cap)
    desc = OxmlElement("w:tblDescription"); desc.set(qn("w:val"), "Revenue rose each quarter."); pr.append(desc)
    p = tmp_path / "t.docx"; d.save(p)
    text, acc = process_office(str(p), "docx")
    assert not [i for i in acc["issues"] if i["code"].startswith("TABLE_")]
    assert "Table: Budget by quarter Revenue rose each quarter." in text and "r1c1" not in text


def test_structured_table_mode_restores_header_checks(sample_docx, monkeypatch):
    monkeypatch.setenv("WEBREADER_TABLES", "structured")
    _, acc = process_office(sample_docx, "docx")
    got = {i["code"] for i in acc["issues"]}
    assert "TABLE_NO_HEADER" in got and "TABLE_DESCRIPTION_MISSING" not in got


def test_docx_marked_header_row_and_title_clear_issues(tmp_path, monkeypatch):
    monkeypatch.setenv("WEBREADER_TABLES", "structured")
    d = docx.Document()
    d.core_properties.title = "Budget"
    d.core_properties.language = "en-US"
    d.add_heading("Budget", 1)
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "Item"
    from docx.oxml import OxmlElement
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))
    p = tmp_path / "ok.docx"; d.save(p)
    _, acc = process_office(str(p), "docx")
    got = {i["code"] for i in acc["issues"]}
    assert "TABLE_NO_HEADER" not in got and "DOC_TITLE_MISSING" not in got and "DOC_LANGUAGE_MISSING" not in got


def test_docx_decorative_figure_is_not_flagged(tmp_path):
    d = docx.Document()
    pic = d.add_picture(_png())
    from lxml import etree
    ns = "http://schemas.microsoft.com/office/drawing/2017/decorative"
    pic._inline.docPr.append(etree.SubElement(pic._inline.docPr, f"{{{ns}}}decorative", val="1"))
    p = tmp_path / "d.docx"; d.save(p)
    _, acc = process_office(str(p), "docx")
    assert not [i for i in acc["issues"] if i["code"].startswith("FIG_")]


# ---------------------------------------------------------------- PPTX

@pytest.fixture
def sample_pptx(tmp_path):
    prs = pptx.Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])           # title slide
    s1.shapes.title.text = "Budget"
    s1.placeholders[1].text = "FY2027"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])           # blank: no title
    tb = s2.shapes.add_textbox(0, 0, 2_000_000, 500_000); tb.text_frame.text = "Orphan text"
    pic = s2.shapes.add_picture(_png(), 0, 1_000_000)          # python-pptx sets alt text to the filename
    pic3 = s2.shapes.add_picture(_png(), 2_000_000, 1_000_000)  # alt text emptied -> missing
    pic3._element.xpath(".//*[local-name()='cNvPr']")[0].set("descr", "")
    pic2 = s2.shapes.add_picture(_png(), 1_000_000, 1_000_000)
    pic2._element.xpath(".//*[local-name()='cNvPr']")[0].set("descr", "A navy square")
    s3 = prs.slides.add_slide(prs.slide_layouts[5])           # title only
    s3.shapes.title.text = "Budget"                            # duplicate title
    gf = s3.shapes.add_table(2, 2, 0, 1_500_000, 3_000_000, 800_000)
    gf.table.first_row = False
    for r in range(2):
        for c in range(2):
            gf.table.cell(r, c).text = f"r{r}c{c}"
    # title not first in reading order: move title to the end of the shape tree
    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    s4.shapes.title.text = "Late title"
    s4.shapes.add_textbox(0, 2_000_000, 1_000_000, 300_000).text_frame.text = "body first"
    el = s4.shapes.title._element; tree = el.getparent(); tree.remove(el); tree.append(el)
    s4.notes_slide.notes_text_frame.text = "Speaker notes"
    path = tmp_path / "p.pptx"; prs.save(path)
    return str(path)


def test_pptx_structure(sample_pptx):
    s = extract_pptx(sample_pptx)
    titles = [b.text for b in s.blocks if b.kind == "slide_title"]
    assert titles == ["Budget", "Budget", "Late title"]
    assert [b.alt_text for b in s.figures()] == ["image.png", "", "A navy square"]
    t = s.tables()[0]
    assert t.rows == [["r0c0", "r0c1"], ["r1c0", "r1c1"]] and t.header_row is False
    assert any(b.kind == "note" and b.text == "Speaker notes" for b in s.blocks)


def test_pptx_issues(sample_pptx):
    _, acc = process_office(sample_pptx, "pptx")
    by = {}
    for i in acc["issues"]:
        by.setdefault(i["code"], []).append(i)
    assert [i["location"] for i in by["SLIDE_TITLE_MISSING"]] == ["Slide 2"]
    assert [i["location"] for i in by["SLIDE_TITLE_DUPLICATE"]] == ["Slide 3"]
    assert [i["location"] for i in by["SLIDE_TITLE_NOT_FIRST"]] == ["Slide 4"]
    assert [i["location"] for i in by["FIG_ALT_MISSING"]] == ["Slide 2"]
    assert [i["location"] for i in by["FIG_ALT_PLACEHOLDER"]] == ["Slide 2"]  # the "image.png" one
    assert [i["location"] for i in by["TABLE_DESCRIPTION_MISSING"]] == ["Slide 3"]
    assert "TABLE_NO_HEADER" not in by


def test_pptx_reading_text(sample_pptx):
    text, _ = process_office(sample_pptx, "pptx")
    assert text.index("Budget") < text.index("Orphan text") < text.index("Image: A navy square")
    assert "Image: image.png" in text  # present alt text is read as authored, even though it is flagged
    assert "Table." in text


# ---------------------------------------------------------------- checks in isolation

def test_alt_text_quality_rules():
    from backend.structure import Block, DocumentStructure, FIGURE
    s = DocumentStructure(format="docx", title="t", language="en", blocks=[
        Block(kind=FIGURE, alt_text="Picture 3", location="a"),
        Block(kind=FIGURE, alt_text="x" * 300, location="b"),
        Block(kind=FIGURE, alt_text="Image of a dog", location="c"),
        Block(kind=FIGURE, alt_text="A dog on a porch", location="d"),
        Block(kind=FIGURE, alt_text="", decorative=True, location="e"),
    ])
    got = {(i.location, i.code) for i in a11y_checks.analyze(s)}
    assert got == {("a", "FIG_ALT_PLACEHOLDER"), ("b", "FIG_ALT_LONG"), ("c", "FIG_ALT_REDUNDANT")}


# ---------------------------------------------------------------- PDF facts

def test_pdf_untagged_facts(tmp_path):
    from backend.tests.test_api import _tiny_pdf
    p = tmp_path / "t.pdf"; p.write_bytes(_tiny_pdf("hello world"))
    acc = process_pdf_accessibility(str(p))
    assert acc["facts"]["has_text_layer"] is True and acc["facts"]["tagged"] is False
    assert {i["code"] for i in acc["issues"]} == {"PDF_UNTAGGED", "PDF_LANGUAGE_MISSING", "PDF_TITLE_MISSING"}
