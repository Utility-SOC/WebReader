"""
PPTX -> DocumentStructure.

Screen readers read a slide's shapes in the order they appear in the file
(the XML shape tree: first = read first = furthest back), which is the
REVERSE of what PowerPoint's Selection Pane shows. We keep that order, record
which shape is the slide title, and record the visual position of each shape so
the checks can tell when reading order and appearance disagree.
"""

from typing import Iterator, List, Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from .structure import (Block, DocumentStructure, SLIDE_TITLE, PARAGRAPH, LIST_ITEM, TABLE, FIGURE, NOTE)

_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}


def _iter_shapes(shapes) -> Iterator:
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def _cnvpr(shape):
    found = shape._element.xpath(".//*[local-name()='cNvPr']")
    return found[0] if found else None


def _alt(shape) -> (str, str, bool):
    c = _cnvpr(shape)
    if c is None:
        return "", "", False
    decorative = bool(shape._element.xpath(".//*[local-name()='decorative']"))
    return (c.get("descr") or "").strip(), c.get("name") or "", decorative


def _is_title(shape) -> bool:
    try:
        return shape.is_placeholder and shape.placeholder_format.type in _TITLE_TYPES
    except Exception:
        return False


def _position(shape):
    try:
        return int(shape.top or 0), int(shape.left or 0)
    except Exception:
        return 0, 0


def extract_pptx(path: str) -> DocumentStructure:
    prs = Presentation(path)
    cp = prs.core_properties
    s = DocumentStructure(
        format="pptx", title=(cp.title or "").strip(), language="",
        metadata={"author": cp.author or "", "created": str(cp.created or ""), "modified": str(cp.modified or ""),
                  "subject": cp.subject or "", "keywords": cp.keywords or "", "slide_count": len(prs.slides),
                  "slide_height": int(prs.slide_height or 0)})

    for i, slide in enumerate(prs.slides, 1):
        loc = f"Slide {i}"
        hidden = slide._element.get("show") == "0"
        has_title_block = False
        order: List[dict] = []  # (xml index, top, left, kind) for the reading-order check

        for idx, shape in enumerate(_iter_shapes(slide.shapes)):
            top, left = _position(shape)

            if getattr(shape, "has_table", False) and shape.has_table:
                tbl = shape.table
                rows = [[c.text.strip() for c in r.cells] for r in tbl.rows]
                alt, name, dec = _alt(shape)
                s.blocks.append(Block(kind=TABLE, rows=rows, header_row=bool(tbl.first_row), location=loc,
                                      alt_text=alt, name=name, extra={"slide": i, "hidden": hidden}))
                order.append({"i": idx, "top": top, "left": left, "kind": "table"})
                continue

            if getattr(shape, "has_chart", False) and shape.has_chart:
                alt, name, dec = _alt(shape)
                s.blocks.append(Block(kind=FIGURE, figure_type="chart", name=name, alt_text=alt, decorative=dec,
                                      location=loc, extra={"slide": i, "hidden": hidden}))
                order.append({"i": idx, "top": top, "left": left, "kind": "figure"})
                continue

            if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.MEDIA, MSO_SHAPE_TYPE.LINKED_PICTURE) \
                    or (shape._element.tag.endswith("}graphicFrame") and shape._element.xpath(".//*[local-name()='relIds']")):
                alt, name, dec = _alt(shape)
                ftype = "media" if shape.shape_type == MSO_SHAPE_TYPE.MEDIA else (
                    "diagram" if shape._element.tag.endswith("}graphicFrame") else "picture")
                s.blocks.append(Block(kind=FIGURE, figure_type=ftype, name=name, alt_text=alt, decorative=dec,
                                      location=loc, extra={"slide": i, "hidden": hidden}))
                order.append({"i": idx, "top": top, "left": left, "kind": "figure"})
                continue

            if shape.has_text_frame and shape.text_frame.text.strip():
                if _is_title(shape):
                    has_title_block = True
                    s.blocks.append(Block(kind=SLIDE_TITLE, text=shape.text_frame.text.strip(), level=1, location=loc,
                                          extra={"slide": i, "hidden": hidden}))
                    order.append({"i": idx, "top": top, "left": left, "kind": "title"})
                else:
                    for p in shape.text_frame.paragraphs:
                        t = "".join(r.text for r in p.runs).strip() or p.text.strip()
                        if not t:
                            continue
                        kind = LIST_ITEM if p.level > 0 or _has_bullet(p) else PARAGRAPH
                        s.blocks.append(Block(kind=kind, text=t, level=p.level, location=loc,
                                              extra={"slide": i, "hidden": hidden}))
                    order.append({"i": idx, "top": top, "left": left, "kind": "text"})
                if not s.language:
                    s.language = _lang_of(shape)
            elif _is_title(shape):
                # an empty title placeholder still counts as "has a title shape" but has no text
                s.blocks.append(Block(kind=SLIDE_TITLE, text="", level=1, location=loc,
                                      extra={"slide": i, "hidden": hidden, "empty": True}))
                order.append({"i": idx, "top": top, "left": left, "kind": "title"})

        if slide.has_notes_slide:
            ntxt = (slide.notes_slide.notes_text_frame.text or "").strip() if slide.notes_slide.notes_text_frame else ""
            if ntxt:
                s.blocks.append(Block(kind=NOTE, text=ntxt, location=loc, extra={"slide": i}))
        s.metadata.setdefault("reading_order", {})[str(i)] = order
        s.metadata.setdefault("slides", {})[str(i)] = {"hidden": hidden, "has_title": has_title_block}
    return s


def _has_bullet(p) -> bool:
    pPr = p._p.pPr
    return bool(pPr is not None and pPr.xpath("./*[local-name()='buChar' or local-name()='buAutoNum']"))


def _lang_of(shape) -> str:
    langs = shape._element.xpath(".//*[local-name()='rPr' or local-name()='endParaRPr']/@lang")
    return langs[0] if langs else ""
